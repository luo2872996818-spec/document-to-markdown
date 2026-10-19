# 测试 Excel 转换器
import re
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
import xlwt
from openpyxl import Workbook

from src.modules.conversion.excel import convertExcel
from src.modules.conversion.result import ConversionError


def test_convertXlsxWithMultipleSheets(tmp_path) -> None:
    """验证多个工作表会转换成多个 Markdown 二级标题。"""
    filePath = tmp_path / "销售数据.xlsx"
    workbook = Workbook()
    firstSheet = workbook.active
    firstSheet.title = "第一季度"
    firstSheet.append(["产品", "金额"])
    firstSheet.append(["钢笔", 120])

    secondSheet = workbook.create_sheet("第二季度")
    secondSheet.append(["产品", "金额"])
    secondSheet.append(["笔记本", 260])
    workbook.save(filePath)

    result = convertExcel(filePath)

    assert result.converterName == "excel"
    assert result.metadata["sheetCount"] == 2
    assert "## 第一季度" in result.markdownText
    assert "## 第二季度" in result.markdownText
    assert "| 钢笔 | 120 |" in result.markdownText
    assert "| 笔记本 | 260 |" in result.markdownText

def setWrongWorksheetDimension(filePath) -> None:
    """把工作表尺寸故意改错，模拟部分 Excel 导出的异常元数据。"""
    with ZipFile(filePath) as sourceArchive:
        archiveItems = [
            (item.filename, sourceArchive.read(item.filename))
            for item in sourceArchive.infolist()
        ]

    with ZipFile(filePath, "w", ZIP_DEFLATED) as targetArchive:
        for fileName, fileData in archiveItems:
            if fileName == "xl/worksheets/sheet1.xml":
                xmlText = fileData.decode("utf-8")
                xmlText = re.sub(
                    r'<dimension ref="[^"]+"\s*/>',
                    '<dimension ref="A1"/>',
                    xmlText,
                    count=1,
                )
                fileData = xmlText.encode("utf-8")
            targetArchive.writestr(fileName, fileData)


def test_convertXlsxWithWrongDimension(tmp_path) -> None:
    """验证错误的尺寸元数据不会导致其他列丢失。"""
    filePath = tmp_path / "错误尺寸.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["姓名", "学院", "学号"])
    sheet.append(["小明", "计算机学院", "2026001"])
    workbook.save(filePath)
    workbook.close()
    setWrongWorksheetDimension(filePath)

    result = convertExcel(filePath)

    assert "| 姓名 | 学院 | 学号 |" in result.markdownText
    assert "| 小明 | 计算机学院 | 2026001 |" in result.markdownText

def test_convertXlsxFormulaWithoutCachedValue(tmp_path) -> None:
    """验证没有缓存计算结果的公式不会变成空单元格。"""
    filePath = tmp_path / "公式.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["数量", "单价", "合计"])
    sheet.append([2, 3, "=A2*B2"])
    workbook.save(filePath)
    workbook.close()

    result = convertExcel(filePath)

    assert "| 2 | 3 | =A2*B2 |" in result.markdownText

def test_convertLegacyXls(tmp_path) -> None:
    """验证旧版 XLS 文件可以读取并转换。"""
    filePath = tmp_path / "旧版数据.xls"
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("数据")
    sheet.write(0, 0, "姓名")
    sheet.write(0, 1, "分数")
    sheet.write(1, 0, "小明")
    sheet.write(1, 1, 95)
    workbook.save(str(filePath))

    result = convertExcel(filePath)

    assert "| 姓名 | 分数 |" in result.markdownText
    assert "| 小明 | 95 |" in result.markdownText


def test_convertMacroEnabledXlsm(tmp_path) -> None:
    """验证 XLSM 扩展名可以使用同一套 Excel 读取逻辑。"""
    filePath = tmp_path / "宏工作簿.xlsm"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["项目", "数量"])
    sheet.append(["测试", 3])
    workbook.save(filePath)
    workbook.close()

    result = convertExcel(filePath)

    assert "## 数据" in result.markdownText
    assert "| 测试 | 3 |" in result.markdownText


def test_rejectCorruptExcel(tmp_path) -> None:
    """验证损坏的 Excel 文件会返回明确错误。"""
    filePath = tmp_path / "损坏.xlsx"
    filePath.write_bytes(b"not a real xlsx file")

    with pytest.raises(ConversionError, match="无法读取 Excel"):
        convertExcel(filePath)