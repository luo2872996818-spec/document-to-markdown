# 实现 Excel 工作簿到 Markdown 的转换
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook

from src.modules.conversion.result import ConversionError, ConversionResult
from src.utils.markdown import normalizeMarkdownText, rowsToMarkdown


def resolveXlsxCell(dataValue: Any, formulaValue: Any) -> Any:
    """公式没有缓存结果时保留公式文本。"""
    if (
        dataValue is None
        and isinstance(formulaValue, str)
        and formulaValue.startswith("=")
    ):
        return formulaValue
    return dataValue


def combineSheetRows(
    dataRows: list[tuple[Any, ...]],
    formulaRows: list[tuple[Any, ...]],
) -> list[list[Any]]:
    """合并缓存值和公式值，生成完整的工作表数据。"""
    rowCount = max(len(dataRows), len(formulaRows))
    combinedRows: list[list[Any]] = []

    for rowIndex in range(rowCount):
        dataRow = dataRows[rowIndex] if rowIndex < len(dataRows) else ()
        formulaRow = formulaRows[rowIndex] if rowIndex < len(formulaRows) else ()
        columnCount = max(len(dataRow), len(formulaRow))
        combinedRows.append([
            resolveXlsxCell(
                dataRow[columnIndex] if columnIndex < len(dataRow) else None,
                formulaRow[columnIndex] if columnIndex < len(formulaRow) else None,
            )
            for columnIndex in range(columnCount)
        ])

    return combinedRows


def readXlsxSheets(filePath: Path) -> list[tuple[str, list[list[Any]]]]:
    """读取 XLSX 或 XLSM 中全部工作表的数据和公式。"""
    # 普通读取模式可以避开部分文件错误的工作表尺寸元数据
    dataWorkbook = load_workbook(
        filename=filePath,
        read_only=False,
        data_only=True,
    )
    formulaWorkbook = load_workbook(
        filename=filePath,
        read_only=False,
        data_only=False,
    )

    try:
        sheets: list[tuple[str, list[list[Any]]]] = []
        for sheet in dataWorkbook.worksheets:
            formulaSheet = formulaWorkbook[sheet.title]
            dataRows = list(sheet.iter_rows(values_only=True))
            formulaRows = list(formulaSheet.iter_rows(values_only=True))
            sheets.append((
                sheet.title,
                combineSheetRows(dataRows, formulaRows),
            ))
        return sheets
    finally:
        dataWorkbook.close()
        formulaWorkbook.close()


def formatXlsCell(sheet: Any, rowIndex: int, columnIndex: int, dateMode: int) -> Any:
    """读取 XLS 单元格，并在必要时把日期序列转换为日期对象。"""
    cell = sheet.cell(rowIndex, columnIndex)
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, dateMode)
        except (ValueError, OverflowError):
            # 日期序列无效时保留原始值，避免整个工作簿转换失败
            return cell.value
    return cell.value


def readXlsSheets(filePath: Path) -> list[tuple[str, list[list[Any]]]]:
    """读取旧版 XLS 文件中全部工作表的数据。"""
    workbook = xlrd.open_workbook(str(filePath), on_demand=True)
    try:
        sheets: list[tuple[str, list[list[Any]]]] = []
        for sheet in workbook.sheets():
            rows = [
                [
                    formatXlsCell(sheet, rowIndex, columnIndex, workbook.datemode)
                    for columnIndex in range(sheet.ncols)
                ]
                for rowIndex in range(sheet.nrows)
            ]
            sheets.append((sheet.name, rows))
        return sheets
    finally:
        workbook.release_resources()


def readWorkbookSheets(filePath: Path) -> list[tuple[str, list[list[Any]]]]:
    """根据 Excel 扩展名选择对应的工作簿读取函数。"""
    suffix = filePath.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return readXlsxSheets(filePath)
    if suffix == ".xls":
        return readXlsSheets(filePath)
    raise ConversionError(f"不支持的 Excel 文件类型：{filePath.suffix}")


def sheetsToMarkdown(
    filePath: Path,
    sheets: list[tuple[str, list[list[Any]]]],
) -> tuple[str, list[str]]:
    """把全部工作表转换为 Markdown，并返回空工作表名称。"""
    documentParts = [f"# {filePath.stem}"]
    emptySheets: list[str] = []

    for sheetName, rows in sheets:
        tableMarkdown = rowsToMarkdown(rows)
        if not tableMarkdown:
            emptySheets.append(sheetName)
            continue
        documentParts.append(f"## {sheetName}\n\n{tableMarkdown}")

    return "\n\n".join(documentParts), emptySheets


def convertExcel(filePath: str | Path) -> ConversionResult:
    """把 Excel 工作簿转换为一个 Markdown 文档。"""
    sourcePath = Path(filePath)
    if not sourcePath.is_file():
        raise ConversionError(f"文件不存在：{sourcePath}")

    try:
        sheets = readWorkbookSheets(sourcePath)
    except ConversionError:
        raise
    except Exception as error:
        raise ConversionError(
            f"无法读取 Excel 文件 {sourcePath.name}：{error}"
        ) from error

    if not sheets:
        raise ConversionError(f"Excel 文件中没有工作表：{sourcePath.name}")

    documentMarkdown, emptySheets = sheetsToMarkdown(sourcePath, sheets)
    tableCount = len(sheets) - len(emptySheets)
    if tableCount == 0:
        raise ConversionError(f"Excel 文件中没有可转换的数据：{sourcePath.name}")

    warningMessage = None
    if emptySheets:
        warningMessage = "以下工作表为空，已跳过：" + "、".join(emptySheets)

    return ConversionResult(
        sourcePath=sourcePath,
        markdownText=normalizeMarkdownText(documentMarkdown),
        converterName="excel",
        warningMessage=warningMessage,
        metadata={
            "sheetCount": len(sheets),
            "tableCount": tableCount,
            "emptySheets": emptySheets,
        },
    )