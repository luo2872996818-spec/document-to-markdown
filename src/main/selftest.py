# 检查打包程序中的核心转换依赖是否完整
import tempfile
from pathlib import Path

import pymupdf
from openpyxl import Workbook

from src.modules.conversion.router import convertDocument


def testCsvConversion(directory: Path) -> None:
    """验证打包环境中的 CSV 转换。"""
    filePath = directory / "self-test.csv"
    filePath.write_text("姓名,数量\n测试,3\n", encoding="utf-8")
    result = convertDocument(filePath)
    if "| 测试 | 3 |" not in result.markdownText:
        raise RuntimeError("CSV 转换结果不正确")


def testExcelConversion(directory: Path) -> None:
    """验证打包环境中的 Excel 转换。"""
    filePath = directory / "self-test.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["姓名", "数量"])
    sheet.append(["测试", 3])
    workbook.save(filePath)
    workbook.close()
    result = convertDocument(filePath)
    if "| 测试 | 3 |" not in result.markdownText:
        raise RuntimeError("Excel 转换结果不正确")


def testPdfConversion(directory: Path) -> None:
    """验证打包环境中的电子版 PDF 转换。"""
    filePath = directory / "self-test.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Packaged PDF self test", fontsize=12)
    document.save(filePath)
    document.close()
    result = convertDocument(filePath)
    if "Packaged PDF self test" not in result.markdownText:
        raise RuntimeError("PDF 转换结果不正确")


def runSelfTest() -> int:
    """运行全部自检，全部成功时返回 0。"""
    tests = [
        ("CSV", testCsvConversion),
        ("Excel", testExcelConversion),
        ("PDF", testPdfConversion),
    ]
    failedCount = 0

    with tempfile.TemporaryDirectory(prefix="document-to-markdown-self-test-") as root:
        directory = Path(root)
        for name, testFunction in tests:
            try:
                testFunction(directory)
                print(f"[PASS] {name}")
            except Exception as error:
                failedCount += 1
                print(f"[FAIL] {name}: {error}")

    print(f"Self test finished: {len(tests) - failedCount}/{len(tests)} passed")
    return 0 if failedCount == 0 else 1